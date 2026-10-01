from collections import defaultdict
from datetime import datetime
from difflib import get_close_matches
from pathlib import Path

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

FILE_PATH = Path(r"FPL_Player_Stats.xlsx")
SEASON = "2026-27"
SHEET_NAME = f"FPL {SEASON}"
ENABLE_FOTMOB = True
FPL_URL = "https://fantasy.premierleague.com/api/"

STAT_FIELDS = {
    "Minutes": "minutes", "Goals Scored": "goals_scored", "Assists": "assists",
    "Clean Sheets": "clean_sheets", "Goals Conceded": "goals_conceded",
    "Own Goals": "own_goals", "Penalties Saved": "penalties_saved",
    "Penalties Missed": "penalties_missed", "Yellow Cards": "yellow_cards",
    "Red Cards": "red_cards", "Saves": "saves", "Bonus": "bonus", "BPS": "bps",
    "Total Points": "total_points", "ICT Index": "ict_index",
    "Influence": "influence", "Creativity": "creativity", "Threat": "threat",
    "Form": "form", "Value": "now_cost", "Health": "chance_of_playing_next_round",
}
FOTMOB_STATS = {
    "minutes": "Minutes", "goals": "Goals", "assists": "Assists",
    "expectedGoals": "xG", "defensiveContribution": "Defensive",
}
DIFFICULTY_COLORS = {1: "008744", 2: "8BD48B", 3: "EEEEEE", 4: "FFB3B3", 5: "E96565"}

def fetch_json(session, url):
    response = session.get(url, timeout=30)
    response.raise_for_status()
    return response.json()

def add_fotmob(session, ws, header, player_rows):

    try:
        records = fetch_json(session, "https://www.fotmob.com/api/players/list/league/47")
        if not isinstance(records, list):
            raise ValueError("Unexpected FotMob player-list format")
        name_to_id = {p["name"].strip(): p["id"] for p in records}
    except (requests.RequestException, ValueError, KeyError, TypeError) as exc:
        print(f"FotMob skipped: {exc}. FPL data will still be saved.")
        return

    names = list(name_to_id)
    failures = 0
    for row, name in player_rows:
        matches = get_close_matches(name, names, n=1, cutoff=0.9)
        matched = name if name in name_to_id else (matches[0] if matches else None)
        if matched is None:
            continue
        try:
            data = fetch_json(session, f"https://www.fotmob.com/api/player/{name_to_id[matched]}/stats")
            if not isinstance(data, dict):
                raise ValueError("Unexpected FotMob stats format")
            for stat, label in FOTMOB_STATS.items():
                for entry in data.get(stat, []):
                    gw = entry.get("round")
                    if isinstance(gw, int) and 1 <= gw <= 38:
                        ws.cell(row, header[f"{label}_GW{gw}"], entry.get("value"))
        except (requests.RequestException, ValueError, KeyError, TypeError, AttributeError):
            failures += 1
    if failures:
        print(f"FotMob stats unavailable for {failures} matched players; their cells are blank.")

def update_workbook(file_path=FILE_PATH):
    file_path = Path(file_path)
    with requests.Session() as session:
        bootstrap = fetch_json(session, FPL_URL + "bootstrap-static/")
        players = bootstrap["elements"]
        if not players:
            raise ValueError("FPL returned no players; workbook has not been changed.")

        first_event = min(bootstrap.get("events", []), key=lambda e: e["id"], default=None)
        if first_event and first_event.get("deadline_time"):
            year = datetime.fromisoformat(first_event["deadline_time"].replace("Z", "+00:00")).year
            if year != int(SEASON.split("-")[0]):
                raise ValueError(f"FPL's first deadline is in {year}, but SEASON is {SEASON}. "
                                 "Check SEASON before running again; no workbook changes saved.")

        fixtures = fetch_json(session, FPL_URL + "fixtures/")
        difficulty = defaultdict(lambda: defaultdict(list))
        for fixture in fixtures:
            gw = fixture.get("event")
            if gw is None:
                continue
            for side in ("h", "a"):
                difficulty[fixture[f"team_{side}"]][gw].append(fixture[f"team_{side}_difficulty"])

        if file_path.exists():
            wb = load_workbook(file_path)
            ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.create_sheet(SHEET_NAME)

            if ws.max_row:
                ws.delete_rows(1, ws.max_row)
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = SHEET_NAME

        headers = ["Player", "Team", "Position", *STAT_FIELDS, *[f"GW{gw}" for gw in range(1, 39)]]
        headers += [f"{label}_GW{gw}" for label in FOTMOB_STATS.values() for gw in range(1, 39)]
        ws.append(headers)
        header = {name: col for col, name in enumerate(headers, 1)}
        teams = {team["id"]: team["name"] for team in bootstrap["teams"]}
        positions = {1: "GK", 2: "DEF", 3: "MID", 4: "FWD"}
        player_rows = []
        failed_history = 0

        for row, player in enumerate(players, 2):
            name = f"{player['first_name']} {player['second_name']}"
            player_rows.append((row, name))
            ws.cell(row, header["Player"], name)
            ws.cell(row, header["Team"], teams.get(player["team"], ""))
            ws.cell(row, header["Position"], positions.get(player["element_type"], ""))
            for label, field in STAT_FIELDS.items():
                ws.cell(row, header[label], player.get(field))

            try:
                history = fetch_json(session, FPL_URL + f"element-summary/{player['id']}/")["history"]
            except (requests.RequestException, ValueError, KeyError, TypeError):
                failed_history += 1
                continue

            points = defaultdict(int)
            for match in history:
                points[match["round"]] += match["total_points"]
            for gw in range(1, 39):
                cell = ws.cell(row, header[f"GW{gw}"])
                if gw in points:
                    cell.value = points[gw]
                    cell.fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
                else:
                    ratings = difficulty[player["team"]].get(gw, [])
                    if ratings:

                        cell.value = sum(ratings) / len(ratings)
                        cell.fill = PatternFill(fill_type="solid", fgColor=DIFFICULTY_COLORS[round(cell.value)])

            if (row - 1) % 50 == 0:
                print(f"Updated {row - 1}/{len(players)} FPL players...")

        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="243746")
        ws.freeze_panes = "D2"
        ws.auto_filter.ref = ws.dimensions
        for col, name in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(col)].width = max(12, len(name) + 2)
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 24
        wb.active = wb.index(ws)
        file_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(file_path)
        if failed_history:
            print(f"FPL history unavailable for {failed_history} players; their GW cells are blank.")
        if ENABLE_FOTMOB:
            add_fotmob(session, ws, header, player_rows)
            wb.save(file_path)
        wb.close()
    print(f"Saved {len(players)} players to '{SHEET_NAME}' in {file_path}")

if __name__ == "__main__":
    try:
        update_workbook()
    except PermissionError:
        raise SystemExit("Cannot save the workbook. Close it in Excel and check folder permissions, then run again.")
